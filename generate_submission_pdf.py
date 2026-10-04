"""
Builds the official 5-slide PDF submission for NxtWave Growth Intern Challenge:
AI60_Growth_OS_NxtWave_Growth_Challenge.pdf
and editable PowerPoint:
AI60_Growth_OS_NxtWave_Growth_Challenge.pptx
"""

import os
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PDF = r"e:\nxtwave\AI60_Growth_OS_NxtWave_Growth_Challenge.pdf"
OUTPUT_PPTX = r"e:\nxtwave\AI60_Growth_OS_NxtWave_Growth_Challenge.pptx"
SCREENSHOTS_DIR = r"e:\nxtwave\screenshots"

# Palette: Clean, professional editorial style
C_PAGE_BG = colors.HexColor('#FFFFFF')
C_HEADER_TEXT = colors.HexColor('#0F172A') # Deep dark slate for maximum contrast on white
C_CARD = colors.HexColor('#0F172A')        # Dark navy for content card container
C_BORDER = colors.HexColor('#1E293B')
C_TEXT = colors.HexColor('#F8FAFC')
C_MUTED = colors.HexColor('#94A3B8')
C_BRAND = colors.HexColor('#4F46E5')
C_GREEN = colors.HexColor('#10B981')
C_AMBER = colors.HexColor('#F59E0B')
C_WHITE = colors.HexColor('#FFFFFF')

def make_pdf():
    # Landscape Letter: 11 x 8.5 inches = 792 x 612 pt
    doc = SimpleDocTemplate(
        OUTPUT_PDF,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    # Custom typography
    title_style = ParagraphStyle(
        'SlideTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=C_HEADER_TEXT,
        spaceAfter=3
    )

    subtitle_style = ParagraphStyle(
        'SlideSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#475569'),
        spaceAfter=12
    )

    body_style = ParagraphStyle(
        'SlideBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=C_TEXT
    )

    body_muted = ParagraphStyle(
        'SlideBodyMuted',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=C_MUTED
    )

    card_header = ParagraphStyle(
        'CardHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=C_TEXT,
        spaceAfter=4
    )

    chip_style = ParagraphStyle(
        'ChipText',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=C_BRAND
    )

    chip_amber_style = ParagraphStyle(
        'ChipAmber',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=C_AMBER
    )

    story = []

    def header_block(title, subtitle, tag="AI60 GROWTH OS · NXTWAVE CHALLENGE"):
        return [
            Table([
                [
                    Paragraph(f"<font color='#6366F1'><b>{tag}</b></font>", chip_style),
                    Paragraph("<font color='#8A93A6'>Candidate: <b>R. Vishnu Sathwick Reddy</b> | Target: Class of 2027</font>", ParagraphStyle('HdrR', fontName='Helvetica', fontSize=8, textColor=C_MUTED, alignment=2))
                ]
            ], colWidths=[360, 360], style=[
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ('TOPPADDING', (0,0), (-1,-1), 0),
            ]),
            Paragraph(title, title_style),
            Paragraph(subtitle, subtitle_style),
        ]

    # ==========================================
    # SLIDE 1: THE STUDENT PROBLEM + CORE THESIS
    # ==========================================
    story.extend(header_block(
        "DON'T SELL THE WORKSHOP. SELL THE OUTCOME.",
        "AI60 Growth OS — Growth Intern Challenge | 500 Final-Year Engineering Students (Class of 2027)"
    ))

    # 3-column layout: Target & Problem | Core Thesis Shift | Upfront Value
    col1 = [
        Paragraph("TARGET AUDIENCE", card_header),
        Paragraph("<b>Class of 2027</b> — Final-year B.Tech / B.E. students across Tier-2 & Tier-3 engineering colleges preparing for technical placement drives.", body_style),
        Spacer(1, 8),
        Paragraph("THE UNDERLYING PROBLEM", card_header),
        Paragraph("Final-year students are overwhelmed by passive, generic webinars. They don't value another attendance certificate. <b>Their urgent motivation is verifiable proof of work they can defend in technical placement rounds.</b>", body_style),
        Spacer(1, 8),
        Paragraph("BARRIER TO CONVERSION", ParagraphStyle('SubWarn', fontName='Helvetica-Bold', fontSize=9, textColor=C_AMBER)),
        Paragraph("Asking students to register upfront asks for commitment before demonstrating personal utility, causing steep drop-offs.", body_muted)
    ]

    col2 = [
        Paragraph("THE CORE THESIS", card_header),
        Paragraph("<i>\"Don't ask students to register — give them a reason to want to register.\"</i>", ParagraphStyle('Thesis', fontName='Helvetica-BoldOblique', fontSize=10.5, leading=14, textColor=C_BRAND)),
        Spacer(1, 10),
        Paragraph("FUNNEL TRANSFORMATION", ParagraphStyle('FTrans', fontName='Helvetica-Bold', fontSize=9.5, textColor=C_TEXT)),
        Spacer(1, 4),
        Paragraph("<font color='#EF4444'><b>Traditional Funnel:</b></font><br/>Landing Page &rarr; High-Friction Form &rarr; Passive Webinar", body_style),
        Spacer(1, 6),
        Paragraph("<font color='#10B981'><b>AI60 Growth Funnel:</b></font><br/>Discover &rarr; Personalize &rarr; <b>Project Passport</b> &rarr; 5-Field Registration", body_style),
        Spacer(1, 8),
        Paragraph("<b>Strategic Insight:</b> Value comes before commitment. The Project Passport turns a generic workshop into a customized build sprint.", body_muted)
    ]

    col3 = [
        Paragraph("WHAT STUDENTS GET BEFORE REGISTERING", card_header),
        Paragraph("In under 60 seconds (answering a few quick questions), the student generates their personalized dossier:", body_muted),
        Spacer(1, 6),
        Paragraph("&bull; <b>Personalized AI Project Concept:</b> Mapped to their domain & skill level.", body_style),
        Spacer(1, 4),
        Paragraph("&bull; <b>60-Minute Build Roadmap:</b> 4 timed phases (Setup &rarr; UI &rarr; AI API &rarr; Demo).", body_style),
        Spacer(1, 4),
        Paragraph("&bull; <b>Interview Positioning:</b> Strategic talking points for recruiters.", body_style),
        Spacer(1, 4),
        Paragraph("&bull; <b>Example Resume Bullet:</b> Placement-ready technical bullet point.", body_style),
        Spacer(1, 8),
        Paragraph("<i>Result: Registration friction collapses because the student is claiming a solution they already own.</i>", body_muted)
    ]

    s1_table = Table([[col1, col2, col3]], colWidths=[235, 240, 245])
    s1_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), C_CARD),
        ('BACKGROUND', (1,0), (1,0), C_CARD),
        ('BACKGROUND', (2,0), (2,0), C_CARD),
        ('BOX', (0,0), (0,0), 1, C_BORDER),
        ('BOX', (1,0), (1,0), 1, C_BRAND),
        ('BOX', (2,0), (2,0), 1, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ('TOPPADDING', (0,0), (-1,-1), 12),
        ('BOTTOMPADDING', (0,0), (-1,-1), 12),
    ]))

    story.append(s1_table)
    story.append(PageBreak())

    # ==========================================
    # SLIDE 2: THE ACQUISITION ENGINE (GROWTH LOOP)
    # ==========================================
    # ==========================================
    # SLIDE 2: THE ACQUISITION ENGINE (GROWTH LOOP)
    # ==========================================
    story.extend(header_block(
        "THE GROWTH LOOP — VALUE-FIRST ARCHITECTURE",
        "VALUE &rarr; REGISTRATION &rarr; REFERRAL &rarr; MORE STUDENTS &rarr; WORKSHOP &rarr; PROJECT &rarr; COMPETITION"
    ))

    loop_steps = [
        ("1. VALUE", "Project Passport", "Blueprint + 60m roadmap"),
        ("2. REGISTER", "5-Field Form", "Class of 2027 minimal friction"),
        ("3. REFERRAL", "Verified Loop", "WhatsApp pre-filled invite"),
        ("4. MORE STUDENTS", "Campus League", "College & referral rankings"),
        ("5. WORKSHOP", "Live Sprint", "Build first AI project in 60m"),
        ("6. PROJECT", "Proof of Work", "GitHub code & resume bullet"),
        ("7. COMPETITION", "Human Judged", "Rs.900 prize pool showcase"),
    ]

    loop_cells = []
    for num, title, desc in loop_steps:
        loop_cells.append([
            Paragraph(f"<b>{num}</b>", ParagraphStyle('LpNum', fontName='Helvetica-Bold', fontSize=8, textColor=C_BRAND, alignment=1)),
            Paragraph(f"<b>{title}</b>", ParagraphStyle('LpT', fontName='Helvetica-Bold', fontSize=9, textColor=C_TEXT, alignment=1)),
            Paragraph(desc, ParagraphStyle('LpD', fontName='Helvetica', fontSize=7.5, textColor=C_MUTED, alignment=1))
        ])

    loop_table = Table([loop_cells], colWidths=[102]*7)
    loop_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('BOX', (0,0), (-1,-1), 1, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(loop_table)
    story.append(Spacer(1, 10))

    # Details: The Passport Engine vs Referral & Campus Mechanics
    p2_left = [
        Paragraph("UPFRONT VALUE: THE PROJECT PASSPORT", card_header),
        Paragraph("<i>\"Don't ask students to register — give them a reason to want to register.\"</i>", ParagraphStyle('CoreT', fontName='Helvetica-BoldOblique', fontSize=9, textColor=C_BRAND)),
        Spacer(1, 4),
        Paragraph("Taking 4 diagnostic inputs (Branch, Skill level, Interest, Preferred domain), the engine generates an interview-ready dossier in 60 seconds:", body_style),
        Spacer(1, 4),
        Paragraph("&bull; <b>1. Project Idea & Title:</b> Concrete problem formulation.", body_muted),
        Paragraph("&bull; <b>2. 60-Minute Scope:</b> What they build live in the sprint.", body_muted),
        Paragraph("&bull; <b>3. Suggested Tech Stack:</b> Modern APIs & frameworks.", body_muted),
        Paragraph("&bull; <b>4. 4-Phase Roadmap:</b> 0-15m Setup &rarr; 15-35m UI &rarr; 35-50m AI &rarr; 50-60m Demo.", body_muted),
        Paragraph("&bull; <b>5. Resume Bullet:</b> Placement-ready technical point.", body_muted),
        Paragraph("&bull; <b>6. Interview Talking Point:</b> Architectural defense statement.", body_muted),
        Spacer(1, 6),
        Paragraph("<b>Strategic Insight:</b> Give value before asking for registration. Students claim an asset they already own.", ParagraphStyle('ValP', fontName='Helvetica-Bold', fontSize=8, textColor=C_GREEN))
    ]

    p2_right = [
        Paragraph("DISTRIBUTION & REWARD INCENTIVE SYSTEM", card_header),
        Paragraph("<i>\"Turn registered students & creators into active distribution channels.\"</i>", ParagraphStyle('CoreT2', fontName='Helvetica-BoldOblique', fontSize=8.5, textColor=C_BRAND)),
        Spacer(1, 4),
        Paragraph("<b>1. Student Creator Challenge (Rs.300 Prize)</b>", ParagraphStyle('QCr', fontName='Helvetica-Bold', fontSize=8.5, textColor=colors.HexColor('#EC4899'))),
        Paragraph("Students create promotional reels, WhatsApp creatives, or posts with tracking codes.<br/><b>Scoring Rubric:</b> Qualified Regs 60% + CTR 20% + Engagement 10% + Creativity 10%.", body_muted),
        Spacer(1, 4),
        Paragraph("<b>2. Campus Referral Leaderboard (Rs.500 Prize Pool)</b>", ParagraphStyle('QRef', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_GREEN)),
        Paragraph("Prizes: #1 Rahul Rs.250 &bull; #2 Priya Rs.150 &bull; #3 Arjun Rs.100.<br/><b>Anti-Fraud Rule:</b> Only verified unique student registrations count. Duplicate emails blocked.", body_muted),
        Spacer(1, 4),
        Paragraph("<b>3. AI Project Competition (Rs.900 Prize Pool)</b>", ParagraphStyle('QCamp', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_BRAND)),
        Paragraph("Prizes: #1 Rs.400 &bull; #2 Rs.300 &bull; #3 Rs.200 for best project code submitted post-workshop.<br/><b>5-Point Rubric:</b> Functionality 30%, Creativity 25%, AI 20%, UX 15%, Defense 10%.", body_muted),
    ]

    p2_table = Table([[p2_left, p2_right]], colWidths=[355, 365])
    p2_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('BOX', (0,0), (-1,-1), 1, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(p2_table)
    story.append(PageBreak())

    # ==========================================
    # SLIDE 3: HOW WE REACH 500 (PLANNING MODEL)
    # ==========================================
    story.extend(header_block(
        "500-REGISTRATION PLANNING MODEL & Rs.2,000 BUDGET",
        "PLANNING ASSUMPTION — NOT ACTUAL CAMPAIGN RESULTS (Hypotheses to validate during 7-day experiment)"
    ))

    # Banner for data honesty
    story.append(Table([[
        Paragraph("<b>NOTE:</b> 500-registration target model — planning assumptions, not campaign results. Levers model channel distribution, not guaranteed acquisition.", chip_amber_style)
    ]], colWidths=[720], style=[
        ('BACKGROUND', (0,0), (0,0), colors.HexColor('#221D12')),
        ('BOX', (0,0), (0,0), 1, C_AMBER),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('TOPPADDING', (0,0), (0,0), 4),
        ('BOTTOMPADDING', (0,0), (0,0), 4),
    ]))
    story.append(Spacer(1, 8))

    # Left: Channel mix breakdown table | Right: Rs.2,000 Budget Table
    p3_left = [
        Paragraph("500-REGISTRATION CHANNEL MIX (HYPOTHESIS)", card_header),
        Table([
            [Paragraph("<b>Channel</b>", ParagraphStyle('Th', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT)),
             Paragraph("<b>Strategic Role</b>", ParagraphStyle('Th', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT)),
             Paragraph("<b>Regs</b>", ParagraphStyle('Th', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT, alignment=2)),
             Paragraph("<b>Share</b>", ParagraphStyle('Th', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT, alignment=2))],
            [Paragraph("Campus Captains + Chapters", body_style), Paragraph("25 captains &times; 8 regs avg (planning assumption)", body_muted), Paragraph("<b>200</b>", body_style), Paragraph("40%", body_muted)],
            [Paragraph("Student Referral Engine", body_style), Paragraph("Verified unique peer-to-peer invites", body_muted), Paragraph("<b>150</b>", body_style), Paragraph("30%", body_muted)],
            [Paragraph("Student Creator Challenge", body_style), Paragraph("Student promotional creators (Rs.300 prize)", body_muted), Paragraph("<b>50</b>", body_style), Paragraph("10%", body_muted)],
            [Paragraph("Organic Social + Communities", body_style), Paragraph("Placement WhatsApp, Telegram, LinkedIn", body_muted), Paragraph("<b>100</b>", body_style), Paragraph("20%", body_muted)],
            [Paragraph("<b>TOTAL TARGET</b>", ParagraphStyle('Tot', fontName='Helvetica-Bold', fontSize=9, textColor=C_BRAND)),
             Paragraph("<b>Planning Baseline</b>", ParagraphStyle('Tot', fontName='Helvetica-Bold', fontSize=8, textColor=C_BRAND)),
             Paragraph("<b>500</b>", ParagraphStyle('Tot', fontName='Helvetica-Bold', fontSize=9, textColor=C_BRAND, alignment=2)),
             Paragraph("<b>100%</b>", ParagraphStyle('Tot', fontName='Helvetica-Bold', fontSize=8, textColor=C_BRAND, alignment=2))]
        ], colWidths=[120, 130, 45, 45], style=[
            ('LINEBELOW', (0,0), (-1,0), 1, C_BORDER),
            ('LINEBELOW', (0,-1), (-1,-1), 1, C_BORDER),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]),
        Spacer(1, 6),
        Paragraph("<b>Data Integrity Note:</b> 500-registration target model — planning assumptions, not campaign results. Levers model channel distribution, not guaranteed acquisition.", body_muted)
    ]

    p3_right = [
        Paragraph("TOTAL BUDGET ALLOCATION: EXACTLY Rs.2,000", card_header),
        Paragraph("<i>Strict financial discipline. Zero budget leak.</i>", ParagraphStyle('SpL', fontName='Helvetica-Oblique', fontSize=8.5, textColor=C_MUTED)),
        Spacer(1, 6),
        Table([
            [Paragraph("<b>Allocation</b>", ParagraphStyle('Th2', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT)),
             Paragraph("<b>Strategic Purpose</b>", ParagraphStyle('Th2', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT)),
             Paragraph("<b>Amount</b>", ParagraphStyle('Th2', fontName='Helvetica-Bold', fontSize=8, textColor=C_TEXT, alignment=2))],
            [Paragraph("<font color='#EC4899'><b>Creator Challenge</b></font>", body_style), Paragraph("Rs.300 prize for top promotional creator (qualified regs + score)", body_muted), Paragraph("<b>Rs.300</b>", body_style)],
            [Paragraph("<font color='#10B981'><b>Referral Rewards</b></font>", body_style), Paragraph("Top 3 verified referrers (Rs.250 / Rs.150 / Rs.100)", body_muted), Paragraph("<b>Rs.500</b>", body_style)],
            [Paragraph("<font color='#818CF8'><b>AI Project Competition</b></font>", body_style), Paragraph("Post-workshop awards (Rs.400 / Rs.300 / Rs.200) — post-workshop incentive", body_muted), Paragraph("<b>Rs.900</b>", body_style)],
            [Paragraph("<font color='#F59E0B'><b>Contingency</b></font>", body_style), Paragraph("Unallocated reserve buffer for campaign safety", body_muted), Paragraph("<b>Rs.300</b>", body_style)],
            [Paragraph("<b>TOTAL</b>", ParagraphStyle('TotB', fontName='Helvetica-Bold', fontSize=9, textColor=C_BRAND)),
             Paragraph("<b>Hard budget ceiling</b>", ParagraphStyle('TotB', fontName='Helvetica-Bold', fontSize=8, textColor=C_BRAND)),
             Paragraph("<b>Rs.2,000</b>", ParagraphStyle('TotB', fontName='Helvetica-Bold', fontSize=9, textColor=C_GREEN, alignment=2))]
        ], colWidths=[125, 170, 50], style=[
            ('LINEBELOW', (0,0), (-1,0), 1, C_BORDER),
            ('LINEBELOW', (0,-1), (-1,-1), 1, C_BORDER),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]),
        Spacer(1, 6),
        Paragraph("<b>Post-Workshop Note:</b> The Rs.900 project competition is a post-workshop engagement incentive (not acquisition spend). The Rs.300 creator challenge funds peer acquisition directly.", body_muted)
    ]

    p3_table = Table([[p3_left, p3_right]], colWidths=[355, 365])
    p3_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('BOX', (0,0), (-1,-1), 1, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(p3_table)
    story.append(PageBreak())

    # ==========================================
    # SLIDE 4: I BUILT THE GROWTH SYSTEM
    # ==========================================
    story.extend(header_block(
        "I DIDN'T JUST WRITE THE PLAN. I BUILT THE LOOP.",
        "Working prototype with full-funnel instrumentation, attribution, and experimentation framework."
    ))

    # Embed 4 cropped screenshots from the active build
    img_w, img_h = 170, 105
    im_hero = Image(os.path.join(SCREENSHOTS_DIR, "hero_crop.png"), width=img_w, height=img_h)
    im_dash = Image(os.path.join(SCREENSHOTS_DIR, "dashboard_crop.png"), width=img_w, height=img_h)
    im_camp = Image(os.path.join(SCREENSHOTS_DIR, "campus_crop.png"), width=img_w, height=img_h)
    im_sim = Image(os.path.join(SCREENSHOTS_DIR, "simulator_crop.png"), width=img_w, height=img_h)

    sc_row = [
        [im_hero, im_dash, im_camp, im_sim],
        [
            Paragraph("<b>1. PROJECT PASSPORT</b><br/><font color='#8A93A6'>Immediate personalized utility; 60m roadmap & resume bullet.</font>", body_style),
            Paragraph("<b>2. GROWTH DASHBOARD</b><br/><font color='#8A93A6'>Strict 3-tier integrity: Actual vs. Simulation vs. Planning.</font>", body_style),
            Paragraph("<b>3. CAMPUS LEAGUE</b><br/><font color='#8A93A6'>Inter-college rivalry driving peer-to-peer distribution.</font>", body_style),
            Paragraph("<b>4. SCENARIO PLANNER</b><br/><font color='#8A93A6'>Live sensitivity testing against the Rs.2,000 budget model.</font>", body_style),
        ]
    ]

    sc_table = Table(sc_row, colWidths=[178, 178, 178, 178])
    sc_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('BOX', (0,0), (-1,-1), 1, C_BORDER),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(sc_table)
    story.append(Spacer(1, 10))

    # Architecture summary table
    sys_summary = [
        [Paragraph("<b>Component</b>", card_header), Paragraph("<b>Growth Objective</b>", card_header), Paragraph("<b>Execution Proof</b>", card_header)],
        [Paragraph("<b>AI Project Passport</b>", body_style), Paragraph("Eliminate registration commitment resistance", body_muted), Paragraph("A few quick questions &rarr; instant 60-min project blueprint & resume bullet", body_style)],
        [Paragraph("<b>Registration Flow</b>", body_style), Paragraph("Low-friction capture for Class of 2027", body_muted), Paragraph("5 fields only (Name, Email, College, Branch, Year 2027)", body_style)],
        [Paragraph("<b>Referral & League</b>", body_style), Paragraph("Organic viral distribution via collegiate pride", body_muted), Paragraph("Requires 2 verified friend signups; real-time campus ranks", body_style)],
        [Paragraph("<b>Creator Challenge</b>", body_style), Paragraph("Student-led content acquisition & testbed", body_muted), Paragraph("Rs.300 prize; judged on qualified regs (60%), CTR (20%), score (86 top)", body_style)],
    ]
    sys_table = Table(sys_summary, colWidths=[140, 260, 310])
    sys_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('LINEBELOW', (0,0), (-1,0), 1, C_BORDER),
        ('LINEBELOW', (0,1), (-1,-2), 0.5, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(sys_table)
    story.append(PageBreak())

    # ==========================================
    # SLIDE 5: GROWTH JUDGMENT + AI LEARNING
    # ==========================================
    story.extend(header_block(
        "AI PROPOSES. EXPERIMENTS MEASURE. I DECIDE.",
        "Demonstrating growth judgment, self-awareness, and disciplined execution tradeoffs."
    ))

    j_col1 = [
        Paragraph("WHAT CHANGED?", card_header),
        Paragraph("<i>\"I initially considered a landing-page-first approach. I changed to a value-first Project Passport + referral system because acquiring a student should create an opportunity to acquire the next student.\"</i>", ParagraphStyle('Ans1', fontName='Helvetica-Oblique', fontSize=9, leading=13.5, textColor=C_TEXT)),
        Spacer(1, 10),
        Paragraph("STRATEGIC RATIONALE", ParagraphStyle('HdrSub', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_BRAND)),
        Paragraph("A standard webinar landing page is linear and leaks traffic at the form. By inverting the funnel to deliver the AI Project Passport upfront, each registrant experiences tangible career utility, converting them into active campus referrers.", body_muted),
        Spacer(1, 8),
        Paragraph("<b>Core Loop:</b> Value &rarr; Registration &rarr; Referral &rarr; More Students.", ParagraphStyle('KL', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_GREEN))
    ]

    j_col2 = [
        Paragraph("WHAT DID AI SUGGEST THAT I REJECTED?", card_header),
        Paragraph("<i>\"AI suggested generic paid advertising and aggressive urgency tactics. I rejected these because Rs.2,000 is too small for reliable cold acquisition and fake urgency harms trust. I replaced paid ads with a Rs.300 Student Creator Challenge instead.\"</i>", ParagraphStyle('Ans2', fontName='Helvetica-Oblique', fontSize=9, leading=13.5, textColor=C_TEXT)),
        Spacer(1, 10),
        Paragraph("WHY HUMAN JUDGMENT MATTERS", ParagraphStyle('HdrSub2', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_AMBER)),
        Paragraph("&bull; <b>No Fake Scarcity:</b> Engineering seniors detect and reject fake countdowns immediately.", body_muted),
        Paragraph("&bull; <b>Creator Incentives over Cold Ads:</b> Rs.300 prize motivates student creators with localized collegiate trust.", body_muted),
        Paragraph("&bull; <b>Quality-First Metric:</b> Views & clicks are useful, but decisions (KILL/SCALE) rely on qualified registrations.", body_muted)
    ]

    j_col3 = [
        Paragraph("WHAT WOULD I IMPROVE WITH ANOTHER 24 HOURS?", card_header),
        Paragraph("<i>\"I would pilot the campaign with 3 colleges, run A/B tests on the referral incentive and creator formats, and scale Creator B's portfolio angle across all campus ambassadors.\"</i>", ParagraphStyle('Ans3', fontName='Helvetica-Oblique', fontSize=9, leading=13.5, textColor=C_TEXT)),
        Spacer(1, 10),
        Paragraph("24-HOUR EXECUTION DRILL", ParagraphStyle('HdrSub3', fontName='Helvetica-Bold', fontSize=8.5, textColor=C_BRAND)),
        Paragraph("&bull; <b>3-College Pilot:</b> Test organic WhatsApp broadcasts at VIT, Amrita, and SRM.", body_muted),
        Paragraph("&bull; <b>Scale Creator Winner:</b> Replicate Creator B's 60s prototype format across campus captains.", body_muted),
        Paragraph("&bull; <b>Attendance Optimization:</b> Deploy WhatsApp calendar reminders post-registration.", body_muted),
        Spacer(1, 10),
        Paragraph("<b>AI proposes. Experiments measure. I decide.</b>", ParagraphStyle('EndQuote', fontName='Helvetica-Bold', fontSize=9.5, leading=13, textColor=C_GREEN, alignment=1))
    ]

    j_table = Table([[j_col1, j_col2, j_col3]], colWidths=[235, 240, 245])
    j_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), C_CARD),
        ('BOX', (0,0), (-1,-1), 1, C_BORDER),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(j_table)

    doc.build(story)
    print(f"Successfully generated PDF: {OUTPUT_PDF}")

def make_pptx():
    prs = Presentation()
    # 16:9 widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Dark background color
    bg_rgb = RGBColor(15, 17, 23)
    card_rgb = RGBColor(22, 25, 34)
    border_rgb = RGBColor(35, 40, 56)
    text_rgb = RGBColor(241, 243, 249)
    muted_rgb = RGBColor(138, 147, 166)
    brand_rgb = RGBColor(99, 102, 241)
    green_rgb = RGBColor(16, 185, 129)
    amber_rgb = RGBColor(245, 158, 11)

    slides_data = [
        {
            "num": "01",
            "tag": "AI60 GROWTH OS · NXTWAVE CHALLENGE",
            "title": "DON'T SELL THE WORKSHOP. SELL THE OUTCOME.",
            "subtitle": "AI60 Growth OS — Growth Intern Challenge | 500 Final-Year Engineering Students (Class of 2027)",
            "cards": [
                ("TARGET & THE PROBLEM", [
                    "Target: Class of 2027 final-year engineering students across Tier-2/3 institutions.",
                    "The Problem: Students are bombarded with generic webinars. They don't value another certificate.",
                    "Core Driver: Urgent placement pressure — need tangible proof of work they can defend in interviews.",
                    "Barrier: Upfront registration forms demand personal data before proving relevance, triggering drop-offs."
                ]),
                ("THE CORE THESIS", [
                    "\"Don't ask students to register — give them a reason to want to register.\"",
                    "Traditional: Landing Page -> Form -> Passive Webinar",
                    "AI60 Funnel: Discover -> Personalize -> Project Passport -> Register -> Refer -> Build -> Compete",
                    "Value precedes commitment: The Project Passport turns a generic workshop into a customized build sprint."
                ]),
                ("UPFRONT UTILITY (PROJECT PASSPORT)", [
                    "In 60 seconds (Branch, Skill, Interest, Domain), students get:",
                    "• Tailored AI Project Blueprint matched to placement domain",
                    "• 60-Minute Build Roadmap (Setup -> UI -> AI Logic -> Demo)",
                    "• Recruiter Placement-Ready Resume Bullet",
                    "• Technical Interview Talking Point for placement defense",
                    "Result: High intent to register and attend because they own the project outcome."
                ])
            ]
        },
        {
            "num": "02",
            "tag": "GROWTH LOOP & DISTRIBUTION ENGINE",
            "title": "THE GROWTH LOOP — VALUE-FIRST ARCHITECTURE",
            "subtitle": "VALUE → REGISTRATION → REFERRAL → MORE STUDENTS → WORKSHOP → PROJECT → COMPETITION",
            "cards": [
                ("1. THE PASSPORT ENGINE", [
                    "Diagnostic inputs: Branch + Skill Level + Interest + Preferred Domain.",
                    "Instant dossier: 60-min build roadmap, tech stack, resume bullet & interview point.",
                    "Registration Friction: 5 essential fields pre-filled with passport context.",
                    "Insight: Give value upfront to eliminate commitment resistance."
                ]),
                ("2. QUALITY-GATED REFERRAL & CREATOR LOOP", [
                    "Student Creator Challenge: ₹300 prize (Qual Regs 60%, CTR 20%, Engagement 10%, Creativity 10%).",
                    "Referral Leaderboard: ₹500 rewards (#1 ₹250, #2 ₹150, #3 ₹100).",
                    "Anti-fraud rule: Qualified Referral = unique student + valid registration + verification.",
                    "Share on WhatsApp: Pre-filled peer invite highlighting the personalized build sprint."
                ]),
                ("3. POST-WORKSHOP COMPETITION", [
                    "AI Project Competition: #1 ₹400, #2 ₹300, #3 ₹200 (Total ₹900 prizes).",
                    "5-Point Rubric: Functionality 30%, Creativity 25%, AI usage 20%, UX 15%, Presentation 10%.",
                    "Human-in-the-Loop: AI provides scoring support; final judging is strictly human-controlled.",
                    "Drives registration-to-attendance and builds genuine portfolio projects."
                ])
            ]
        },
        {
            "num": "03",
            "tag": "PLANNING MODEL · NOT ACTUAL CAMPAIGN RESULTS",
            "title": "500 REGISTRATIONS — PLANNING MODEL & ₹2,000 BUDGET",
            "subtitle": "500-registration target model — planning assumptions, not campaign results.",
            "cards": [
                ("500-REGISTRATION CHANNEL MIX", [
                    "• Campus Captains + Chapters: 200 (40%)",
                    "• Student Referral Engine: 150 (30%)",
                    "• Student Creator Challenge: 50 (10%)",
                    "• Organic Social + Communities: 100 (20%)",
                    "TOTAL TARGET = 500 Registrations (100%)",
                    "LABEL: PLANNING ASSUMPTION — NOT ACTUAL CAMPAIGN RESULTS"
                ]),
                ("STRICT ₹2,000 BUDGET ALLOCATION", [
                    "• Creator Challenge: ₹300 (Prize for best promotional creator)",
                    "• Referral Rewards: ₹500 (#1 ₹250, #2 ₹150, #3 ₹100)",
                    "• AI Project Competition: ₹900 (#1 ₹400, #2 ₹300, #3 ₹200)",
                    "• Contingency: ₹300 (safety buffer)",
                    "TOTAL = EXACTLY ₹2,000",
                    "Note: ₹900 is a post-workshop engagement incentive, not acquisition spend."
                ]),
                ("DECISION FRAMEWORK & SENSITIVITY", [
                    "Framework: KILL / ITERATE / SCALE.",
                    "Creator Challenge: Creator A (ITERATE), Creator B (SCALE, ₹300 winner), Creator C (KILL).",
                    "Insight: Views are useful. Clicks are useful. Registrations matter more. Qualified registrations matter most.",
                    "Sensitivity: Fixed ₹300 prize prevents budget blowup if creative underperforms."
                ])
            ]
        },
        {
            "num": "04",
            "tag": "WORKING ASSET DEMONSTRATION",
            "title": "I DIDN'T JUST WRITE THE PLAN. I BUILT THE LOOP.",
            "subtitle": "Working growth prototype demonstrating value upfront, verified referrals, and disciplined experimentation.",
            "cards": [
                ("1. PROJECT PASSPORT & REFERRALS", [
                    "• AI Project Passport: Branch + Domain customization with resume bullet & interview talking point.",
                    "• Dedicated Referral Hub: Unique code (VISHNU26), referral link, live conversion rate, rank #4.",
                    "• 1-Click WhatsApp Share: Formatted viral invite message.",
                    "• Campus Referral Leaderboard: Top 3 rewards with anti-fraud verification note."
                ]),
                ("2. CREATOR CHALLENGE & FUNNEL", [
                    "• Student Creator Challenge: ₹300 prize simulation (Creator B leads with 31 qualified regs, score 86).",
                    "• Metrics tracked: Reach, Clicks, Qualified Regs, CTR, Score, Status.",
                    "• 8-Step Funnel: Discovery → Passport → Registration → Referral → Qualified Reg → Workshop → Project Submission → Competition.",
                    "• Strictly labeled SIMULATION DATA."
                ]),
                ("3. SCENARIO PLANNER & ADMIN", [
                    "• Scenario Planner: Interactive levers (Captains, Referrals, Creator Challenge, Social, Prize budget).",
                    "• Auto-calculates Projected Regs & Remaining Budget with dual risk warnings.",
                    "• Executive Admin Analytics: Acquisition, Sources, Referral quality, Creator performance, Workshop stages."
                ])
            ]
        },
        {
            "num": "05",
            "tag": "GROWTH JUDGMENT & LEARNING",
            "title": "AI PROPOSES. EXPERIMENTS MEASURE. I DECIDE.",
            "subtitle": "Demonstrating growth judgment, self-awareness, and disciplined execution tradeoffs.",
            "cards": [
                ("WHAT CHANGED?", [
                    "\"I initially considered a landing-page-first approach. I changed to a value-first Project Passport + referral system because acquiring a student should create an opportunity to acquire the next student.\"",
                    "Strategic Rationale: A standard webinar page leaks at the form. Delivering utility first turns registrants into active distribution advocates."
                ]),
                ("WHAT DID AI SUGGEST THAT I REJECTED?", [
                    "\"AI suggested generic paid advertising and aggressive urgency tactics. I rejected these because ₹2,000 is too small for reliable mass acquisition and fake urgency harms trust. I replaced paid ads with a ₹300 Student Creator Challenge instead.\"",
                    "Growth Principle: Views are useful. Clicks are useful. Registrations matter more. Qualified registrations matter most."
                ]),
                ("WHAT WOULD I IMPROVE IN 24 HOURS?", [
                    "\"I would pilot the campaign with 3 colleges, run A/B tests on the referral incentive and creator formats, and scale Creator B's portfolio angle across all campus ambassadors.\"",
                    "Next Steps: Test WhatsApp forwards with VIT, Amrita, SRM; scale winning creative.",
                    "End with: AI proposes. Experiments measure. I decide."
                ])
            ]
        }
    ]

    for s_idx, sdata in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = bg_rgb
        bg_shape.line.color.rgb = bg_rgb

        # Top Tag & Header
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"{sdata['tag']}  |  SLIDE {sdata['num']} OF 05"
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = brand_rgb

        p1 = tf.add_paragraph()
        p1.text = sdata['title']
        p1.font.bold = True
        p1.font.size = Pt(22)
        p1.font.color.rgb = text_rgb
        p1.space_before = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = sdata['subtitle']
        p2.font.size = Pt(11)
        p2.font.color.rgb = muted_rgb
        p2.space_before = Pt(2)

        # 3 Card Columns
        card_w = Inches(3.64)
        card_h = Inches(5.1)
        gap = Inches(0.39)
        top_y = Inches(1.8)

        for c_idx, (card_title, bullets) in enumerate(sdata['cards']):
            left_x = Inches(0.8) + c_idx * (card_w + gap)

            card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, card_w, card_h)
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = card_rgb
            card_shape.line.color.rgb = border_rgb
            card_shape.line.width = Pt(1)

            ctb = slide.shapes.add_textbox(left_x + Inches(0.2), top_y + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
            ctf = ctb.text_frame
            ctf.word_wrap = True
            ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0

            cp0 = ctf.paragraphs[0]
            cp0.text = card_title
            cp0.font.bold = True
            cp0.font.size = Pt(13)
            cp0.font.color.rgb = text_rgb
            cp0.space_after = Pt(10)

            for b in bullets:
                bp = ctf.add_paragraph()
                bp.text = b
                bp.font.size = Pt(10)
                bp.font.color.rgb = muted_rgb if not b.startswith("•") and not b.startswith("\"") and not "TOTAL" in b else text_rgb
                bp.space_after = Pt(6)

    prs.save(OUTPUT_PPTX)
    print(f"Successfully generated PPTX: {OUTPUT_PPTX}")

if __name__ == "__main__":
    make_pdf()
    make_pptx()
